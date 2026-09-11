"""Independently reconcile source data, causal decisions and the delivered workbook.

Run with the same numpy/pandas/scipy/openpyxl environment as solve_q2.py.
Reads existing outputs without overwriting them. --replay re-solves all day-ahead
plans and all January parameter candidates to verify numerical reproducibility.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from solve_q2 import plan_day, execute_day

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '问题2/结果'
TOL = 1e-6


def close(a, b):
    np.testing.assert_allclose(a, b, rtol=0, atol=TOL)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--replay', action='store_true')
    args = parser.parse_args()
    f = pd.read_csv(OUT/'problem2_dispatch_all_2025.csv')
    summary = json.loads((OUT/'problem2_summary.json').read_text(encoding='utf-8'))
    daily = pd.read_csv(OUT/'problem2_daily_summary.csv')
    audits = pd.read_csv(OUT/'problem2_information_audit.csv')
    source = ROOT/'公共代码/数据/master_10min.csv'
    assert hashlib.sha256(source.read_bytes()).hexdigest() == summary['master_sha256']
    dates = pd.date_range('2025-01-01', '2025-12-31').strftime('%Y-%m-%d').tolist()
    assert f.date.drop_duplicates().tolist() == dates
    assert len(f) == 365*144 and not f.duplicated(['date', 'slot']).any()
    close(f.slot, np.tile(np.arange(144), 365))
    assert np.isfinite(f.select_dtypes(include='number')).all().all()
    price = pd.read_excel(ROOT/'题目/附件/附件1.xlsx').iloc[:, 1].to_numpy(float)
    for sheet, column in [('小区负载', 'load_actual_kw'), ('光伏发电实际功率', 'pv_actual_kw')]:
        raw = pd.read_excel(ROOT/'题目/附件/附件2.xlsx', sheet_name=sheet)
        assert pd.to_datetime(raw.iloc[:, 0]).dt.strftime('%Y-%m-%d').tolist() == dates
        close(raw.iloc[:, 1:].to_numpy(float).ravel(), f[column])
    close(f.price_yuan_per_kwh, np.tile(price, 365))
    net = (f.load_actual_kw-f.pv_actual_kw).to_numpy().reshape(365, 144)
    close(f.soc_start_kwh.iloc[0], 6000)
    close(f.soc_start_kwh.to_numpy()[1:], f.soc_end_kwh.to_numpy()[:-1])
    close(f.soc_end_kwh-f.soc_start_kwh, .9*f.charge_kw/6-f.discharge_kw/5.4)
    close(f.plan_kw+f.emergency_kw+f.pv_actual_kw+f.discharge_kw,
          f.load_actual_kw+f.charge_kw+f.surplus_kw)
    for stem in ['plan', 'charge', 'discharge', 'emergency']:
        assert f[stem+'_kw'].min() >= -TOL
        close(f[stem+'_kwh'], f[stem+'_kw']/6)
    assert f.surplus_kw.min() >= -TOL
    assert f[['soc_start_kwh', 'soc_end_kwh']].min().min() >= 1200-TOL
    assert f[['soc_start_kwh', 'soc_end_kwh']].max().max() <= 10800+TOL
    assert f[['charge_kw', 'discharge_kw']].max().max() <= 5000+TOL
    assert not ((f.charge_kw > TOL) & (f.discharge_kw > TOL)).any()
    close(f.plan_cost_yuan, f.price_yuan_per_kwh*f.plan_kwh)
    close(f.emergency_cost_yuan, 5*f.price_yuan_per_kwh*f.emergency_kwh)
    close(f.total_cost_yuan, f.plan_cost_yuan+f.emergency_cost_yuan)
    # Rebuild every prediction directly from a strict history prefix.
    for i, (date, g) in enumerate(f.groupby('date', sort=True)):
        w, alpha = (14, .8) if i < 31 else (summary['selected_window_days'], summary['selected_quantile'])
        pred = np.quantile(net[max(0, i-w):i], alpha, axis=0) if i else np.zeros(144)
        close(pred, g.net_forecast_kw)
        if i:
            assert audits.history_through_date.iloc[i] == dates[i-1]
            close(g.planned_soc_end_kwh.iloc[-1], 6000)
        close(g.planned_soc_start_kwh.iloc[0], g.soc_start_kwh.iloc[0])
        close(g.planned_soc_start_kwh.to_numpy()[1:], g.planned_soc_end_kwh.to_numpy()[:-1])
        close(g.planned_soc_end_kwh-g.planned_soc_start_kwh,
              .9*g.planned_charge_kw/6-g.planned_discharge_kw/5.4)
        assert not ((g.planned_charge_kw>TOL)&(g.planned_discharge_kw>TOL)).any()
        assert g[['planned_charge_kw', 'planned_discharge_kw']].min().min() >= -TOL
        assert g[['planned_charge_kw', 'planned_discharge_kw']].max().max() <= 5000+TOL
        assert g[['planned_soc_start_kwh', 'planned_soc_end_kwh']].min().min() >= 1200-TOL
        assert g[['planned_soc_start_kwh', 'planned_soc_end_kwh']].max().max() <= 10800+TOL
        assert (g.plan_kw+g.planned_discharge_kw-g.planned_charge_kw-pred).min() >= -TOL
        ex = execute_day(g.plan_kw.to_numpy(), net[i], g.soc_start_kwh.iloc[0], idle_storage=i==0)
        for key in ['charge', 'discharge', 'emergency', 'surplus']:
            close(ex[key], g[key+'_kw'])
        if args.replay and i:
            plan = plan_day(pred, price, g.soc_start_kwh.iloc[0])
            close(plan['cost'], g.plan_cost_yuan.sum())
        if i % 60 == 0:
            print(f'verified through {date}', flush=True)
    delivered = f[f.date >= '2025-02-01']
    assert len(delivered) == 334*144
    for col in ['plan_kwh','emergency_kwh','charge_kwh','discharge_kwh',
                'plan_cost_yuan','emergency_cost_yuan','total_cost_yuan']:
        close(delivered[col].sum(), summary[col])
        close(delivered.groupby('date')[col].sum().to_numpy(), daily[col])
    candidates = pd.read_csv(OUT/'problem2_january_validation.csv')
    best = candidates.sort_values(['score_yuan','window_days','quantile']).iloc[0]
    close(best.window_days, summary['selected_window_days'])
    close(best['quantile'], summary['selected_quantile'])
    if args.replay:
        initial = f[f.date=='2025-01-15'].soc_start_kwh.iloc[0]
        for row in candidates.itertuples(index=False):
            soc, cost, emergency = initial, 0., 0.
            for i in range(14, 31):
                pred = np.quantile(net[max(0,i-row.window_days):i], row.quantile, axis=0)
                plan = plan_day(pred, price, soc)
                ex = execute_day(plan['purchase'], net[i], soc)
                cost += plan['cost'] + 5*float(price@ex['emergency']/6)
                emergency += ex['emergency'].sum()/6
                soc = ex['soc'][-1]
            close([cost, soc, emergency], [row.validation_cost_yuan, row.terminal_soc_kwh, row.emergency_kwh])
            close(cost+(6000-soc)*price.min()/.9, row.score_yuan)
        print('all 15 January candidates reproduced', flush=True)
    # Reconstruct merged emergency intervals without relying on the export payload.
    expected_events = []
    for date, g in delivered.groupby('date'):
        positive = g.emergency_kwh.to_numpy() > 1e-8
        edges = np.diff(np.r_[False, positive, False].astype(int))
        for a, b in zip(np.where(edges==1)[0], np.where(edges==-1)[0]):
            label = f'{a//6:02d}:{(a%6)*10:02d}-{b//6:02d}:{(b%6)*10:02d}'
            expected_events.append((date,label,float(g.emergency_kwh.iloc[a:b].sum())))
    events = pd.read_csv(OUT/'problem2_emergency_events.csv')
    assert list(zip(events.date, events.interval)) == [(d,t) for d,t,_ in expected_events]
    close(events.emergency_kwh, [v for _,_,v in expected_events])
    book = load_workbook(OUT/'result2.xlsx', data_only=True)
    template = load_workbook(ROOT/'题目/附件/附件5/result2.xlsx', data_only=True)
    assert book.sheetnames == template.sheetnames
    plan = book['计划购电量']
    assert plan.max_row == 335
    assert [plan.cell(1,c).value for c in range(1,148)] == [template['计划购电量'].cell(1,c).value for c in range(1,148)]
    storage = book['充放电量']
    assert storage.max_row == 2005
    for i,(date,g) in enumerate(delivered.groupby('date')):
        assert pd.Timestamp(plan.cell(i+2,1).value).strftime('%Y-%m-%d') == date
        close([plan.cell(i+2,c).value for c in range(2,146)],g.plan_kwh)
        close([plan.cell(i+2,146).value,plan.cell(i+2,147).value],
              [g.plan_kwh.sum(),g.plan_cost_yuan.sum()])
        r = 2+6*i
        assert pd.Timestamp(storage.cell(r,1).value).strftime('%Y-%m-%d') == date
        for c,field in [(3,'charge_kwh'),(4,'discharge_kwh')]:
            close([storage.cell(r+b,c).value for b in range(6)],g[field].to_numpy().reshape(6,24).sum(axis=1))
        close([storage.cell(r,6).value,storage.cell(r+1,6).value],
              [g.soc_start_kwh.iloc[0],g.soc_end_kwh.iloc[-1]])
    actual_events, last_date, displayed_dates = [], None, set()
    for date,label,energy in book['紧急购电量'].iter_rows(min_row=2, max_col=3, values_only=True):
        if date is not None:
            last_date = pd.Timestamp(date).strftime('%Y-%m-%d')
            displayed_dates.add(last_date)
        if label == '无':
            close(energy, 0)
        elif label is not None:
            actual_events.append((last_date,label,energy))
    assert displayed_dates == set(delivered.date)
    assert [(d,t) for d,t,_ in actual_events] == [(d,t) for d,t,_ in expected_events]
    close([v for _,_,v in actual_events], [v for _,_,v in expected_events])
    assert not any(cell.data_type=='e' for sheet in book for row in sheet for cell in row)
    print(json.dumps(dict(verified=True, delivered_days=334, delivered_slots=48096,
                         emergency_events=len(actual_events), replay=args.replay,
                         total_cost_yuan=summary['total_cost_yuan']), indent=2))


if __name__ == '__main__':
    main()
