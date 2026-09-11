"""Independent Q4 reconciliation: source, information cutoff, physics, fees, XLSX."""
import itertools
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from price_forecast import build_prices, january_selection, N
from solve_q4 import q3

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'问题4/结果'
TOL=3e-5

def close(a,b,tol=TOL):
    assert np.allclose(np.asarray(a,dtype=float),np.asarray(b,dtype=float),atol=tol,rtol=0), f'max diff={np.max(np.abs(np.asarray(a)-np.asarray(b)))}'

def test_information(master):
    prices=master.price_volatile_yuan_per_kwh.to_numpy().reshape(365,N)
    dates=master.date.drop_duplicates().tolist()
    before=build_prices(prices,dates)
    altered=prices.copy()
    # Change every unobserved price from Apr 11 noon onward, including that
    # same day's afternoon. This must not change any earlier issued forecast.
    cutoff=100*N+72
    altered.ravel()[cutoff:]*=7
    after=build_prices(altered,dates)
    for name in ('raw','mean','radius'):
        close(getattr(before,name)[:100],getattr(after,name)[:100])
        close(getattr(before,name)[100,:3],getattr(after,name)[100,:3])
    assert before.selected_model==after.selected_model
    january_altered=prices.copy(); january_altered[31:]*=11
    assert january_selection(prices)[0]==january_selection(january_altered)[0]
    for row in before.audit.itertuples():
        assert not row.observations_end or row.observations_end<=row.issue_datetime
        assert not row.latest_scenario_target_end or row.latest_scenario_target_end<=row.issue_datetime
        assert not row.scenario_issue_last_date or row.scenario_issue_last_date<row.issue_datetime[:10]
    # Nonnegative exposure: box max is exactly the upper-price vector.
    center=np.array([.4,.8,1.2]); radius=np.array([.1,.2,.3]); exposure=np.array([10.,20.,30.])
    corners=np.array([center+radius*np.array(sign) for sign in itertools.product((-1,1),repeat=3)])
    close(max(corners@exposure),(center+radius)@exposure)
    close(np.mean(corners@exposure),corners.mean(axis=0)@exposure)
    return before

def main():
    master=pd.read_csv(ROOT/'公共代码/数据/master_10min.csv').sort_values(['date','slot'])
    price_paths=test_information(master)
    q2=pd.read_csv(ROOT/'问题2/结果/problem2_dispatch_all_2025.csv').sort_values(['date','slot'])
    checks={}
    for variant in ('4-2','4-3'):
        f=pd.read_csv(OUT/f'problem{variant}_dispatch_all_2025.csv')
        g=f[f.date>='2025-02-01'].copy()
        w=pd.read_csv(OUT/f'problem{variant}_rolling_windows.csv.gz')
        audit=pd.read_csv(OUT/f'problem{variant}_information_audit.csv',keep_default_na=False)
        summary=json.loads((OUT/f'problem{variant}_summary.json').read_text())
        assert len(f)==365*144 and len(g)==334*144
        assert len(audit)==365*(1 if variant=='4-2' else 4)
        assert np.array_equal(f[['date','slot']],master[['date','slot']])
        for col in ('load_actual_kw','pv_actual_kw'): close(f[col],master[col])
        close(f.price_yuan_per_kwh,master.price_volatile_yuan_per_kwh)
        target=pd.to_datetime(f.date)+pd.to_timedelta((f.slot+1)*10,unit='min')
        assert (pd.to_datetime(f.target_end)==target).all()
        assert (pd.to_datetime(f.issue_datetime)<target).all()
        assert (f.issue_hour==(0 if variant=='4-2' else f.slot//36*6)).all()
        close(f.soc_start_kwh.iloc[0],6000)
        physical=q3.validate(f)
        close(f.plan_cost_yuan,f.plan_kwh*f.price_yuan_per_kwh)
        close(f.adjustment_up_kwh,np.maximum(f.adjusted_purchase_kwh-f.plan_kwh,0))
        close(f.adjustment_down_kwh,np.maximum(f.plan_kwh-f.adjusted_purchase_kwh,0))
        close(f.adjustment_cost_yuan,f.price_yuan_per_kwh*(1.5*f.adjustment_up_kwh+.5*f.adjustment_down_kwh))
        close(f.emergency_cost_yuan,5*f.price_yuan_per_kwh*f.emergency_kwh)
        close(f.total_cost_yuan,f.plan_cost_yuan+f.adjustment_cost_yuan+f.emergency_cost_yuan)
        for col in q3.SUM_COLS: close(g[col].sum(),summary[col],tol=.002)
        close(w.soc_end_kwh-w.soc_start_kwh,.9*w.charge_kwh-w.discharge_kwh/.9)
        assert w.soc_start_kwh.min()>=1200-TOL and w.soc_end_kwh.min()>=1200-TOL
        assert w.soc_start_kwh.max()<=10800+TOL and w.soc_end_kwh.max()<=10800+TOL
        assert max(w.charge_kwh.max(),w.discharge_kwh.max())<=5000/6+TOL
        assert not ((w.charge_kwh>TOL)&(w.discharge_kwh>TOL)).any()
        if variant=='4-2':
            close(f.adjustment_up_kwh,0); close(f.adjustment_down_kwh,0)
            close(f.pv_forecast_kw,q2.pv_forecast_kw)
            close(w.net_forecast_kw,q2.net_forecast_kw)
            close(w.risk_requirement_kw,q2.reserve_requirement_kw)
            slack=w.purchase_kwh+w.discharge_kwh-w.charge_kwh-w.net_forecast_kw/6
            reserve=w.risk_requirement_kw/6-w.risk_proxy_kwh
            assert (w.discharge_kwh+reserve<=5000/6+TOL).all()
            assert (w.soc_start_kwh>=1200+(w.discharge_kwh+reserve)/.9-TOL).all()
            assert (reserve>=-TOL).all() and w.risk_proxy_kwh.min()>=-TOL
        else:
            slack=w.purchase_kwh+w.risk_proxy_kwh+w.discharge_kwh-w.charge_kwh-(w.net_forecast_kw+w.risk_requirement_kw)/6
            hourly=pd.read_csv(ROOT/'公共代码/数据/forecast_attachment3_hourly.csv').set_index(['issue_datetime','target_datetime'])
            whole=w[w.lead_slot%6==0]
            close(whole.pv_forecast_kw,hourly.loc[pd.MultiIndex.from_frame(whole[['issue_datetime','target_end']]),'pv_forecast_kw'])
        assert slack.min()>=-TOL
        issue=pd.to_datetime(w.issue_datetime)
        assert (pd.to_datetime(w.target_end)==issue+pd.to_timedelta(w.lead_slot*10,unit='min')).all()
        days=issue.dt.dayofyear.to_numpy()-1; phase=(issue.dt.hour//6).to_numpy(); lead=w.lead_slot.to_numpy()-1
        close(w.price_mean_yuan_per_kwh,price_paths.mean[days,phase,lead])
        close(w.price_radius_yuan_per_kwh,price_paths.radius[days,phase,lead])
        close(w.price_used_for_optimization,w.price_mean_yuan_per_kwh+w.price_radius_yuan_per_kwh)
        committed=w[w.committed].set_index('target_end').loc[f.target_end]
        assert len(committed)==len(f) and not committed.index.duplicated().any()
        close(committed.purchase_kwh,f.adjusted_purchase_kwh)
        midnight=w[issue.dt.hour==0].set_index('target_end').loc[f.target_end]
        close(midnight.purchase_kwh,f.plan_kwh)
        assert not w[w.provisional_next_day].committed.any()
        for row in audit.itertuples(): assert row.load_history_through_date<row.date
        book=load_workbook(OUT/f'result{variant}.xlsx',data_only=True)
        expected_names=['计划购电量']+(['调整购电量'] if variant=='4-3' else [])+['充放电量','紧急购电量']
        assert book.sheetnames==expected_names
        sheets=[('计划购电量','plan_kwh','plan_cost_yuan')]+([('调整购电量','adjusted_purchase_kwh','adjustment_cost_yuan')] if variant=='4-3' else [])
        for name,col,fee in sheets:
            sheet=book[name]
            assert [sheet.cell(1,j).value for j in range(2,146)]==f.time_interval.iloc[:144].tolist()
            for r,(date,day) in enumerate(g.groupby('date'),start=2):
                assert pd.Timestamp(sheet.cell(r,1).value).strftime('%Y-%m-%d')==date
                close([sheet.cell(r,j).value for j in range(2,146)],day[col])
                close([sheet.cell(r,146).value,sheet.cell(r,147).value],[day[col].sum(),day[fee].sum()])
        for i,(date,day) in enumerate(g.groupby('date')):
            r=2+i*6; sheet=book['充放电量']
            assert pd.Timestamp(sheet.cell(r,1).value).strftime('%Y-%m-%d')==date
            close([sheet.cell(r+j,3).value for j in range(6)],day.charge_kwh.to_numpy().reshape(6,24).sum(axis=1))
            close([sheet.cell(r+j,4).value for j in range(6)],day.discharge_kwh.to_numpy().reshape(6,24).sum(axis=1))
            close([sheet.cell(r,6).value,sheet.cell(r+1,6).value],[day.soc_start_kwh.iloc[0],day.soc_end_kwh.iloc[-1]])
        found=[]; current=None
        for date,span,energy in book['紧急购电量'].iter_rows(min_row=2,max_col=3,values_only=True):
            if date is not None: current=pd.Timestamp(date).strftime('%Y-%m-%d')
            if span and span!='无': found.append((current,span,energy))
        ev=pd.read_csv(OUT/f'problem{variant}_emergency_events.csv')
        assert [(a,b) for a,b,_ in found]==list(zip(ev.date,ev.interval))
        close([e for _,_,e in found],ev.emergency_kwh)
        close(ev.emergency_kwh.sum(),g.emergency_kwh.sum())
        for sheet in book:
            assert not any(c.data_type=='e' for row in sheet for c in row)
        checks[variant]={'status':'passed','formal_days':334,'slots':len(g),'optimization_windows':len(audit),'xlsx_sheets':expected_names,'physics':physical}
    checks['logic_tests']={'future_price_perturbation':'passed','january_only_selection':'passed','scenario_expectation_equivalence':'passed','box_robust_equivalence':'passed','actual_daily_prices_in_xlsx':'passed'}
    q3.save_json(OUT/'problem4_verification.json',checks)
    print(json.dumps(checks,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
