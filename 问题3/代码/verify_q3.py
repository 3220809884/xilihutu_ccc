"""Independent read-only verification of problem-3 sources, trajectories and workbook."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "问题3" / "结果"
TOL = 1e-6
DT = 1 / 6


def close(actual, expected):
    np.testing.assert_allclose(actual, expected, rtol=0, atol=TOL)


def time_label(slot: int) -> str:
    return f"{slot // 6:02d}:{slot % 6 * 10:02d}"


def main() -> None:
    frame = pd.read_csv(OUT / "problem3_dispatch_all_2025.csv")
    delivered = frame[frame.date >= "2025-02-01"].reset_index(drop=True)
    daily = pd.read_csv(OUT / "problem3_daily_summary.csv")
    events = pd.read_csv(OUT / "problem3_emergency_events.csv")
    summary = json.loads((OUT / "problem3_summary.json").read_text(encoding="utf-8"))
    master_path = ROOT / "公共代码/数据/master_10min.csv"
    forecast_path = ROOT / "公共代码/数据/forecast_attachment3_hourly.csv"
    assert hashlib.sha256(master_path.read_bytes()).hexdigest() == summary["master_sha256"]
    assert hashlib.sha256(forecast_path.read_bytes()).hexdigest() == summary["forecast_sha256"]
    assert len(frame) == 365 * 144 and len(delivered) == 334 * 144
    assert not frame.duplicated(["date", "slot"]).any()
    close(frame.slot, np.tile(np.arange(144), 365))

    # Independently compare canonical data against the original attachments.
    price = pd.read_excel(ROOT / "题目/附件/附件1.xlsx").iloc[:, 1].to_numpy(float)
    close(frame.price_yuan_per_kwh, np.tile(price, 365))
    for sheet, column in (("小区负载", "load_actual_kw"), ("光伏发电实际功率", "pv_actual_kw")):
        raw = pd.read_excel(ROOT / "题目/附件/附件2.xlsx", sheet_name=sheet)
        close(frame[column], raw.iloc[:, 1:].to_numpy(float).ravel())
    raw3 = pd.read_excel(ROOT / "题目/附件/附件3.xlsx")
    raw3.iloc[:, 0] = raw3.iloc[:, 0].ffill()
    normalized = pd.read_csv(forecast_path).sort_values(["issue_date", "issue_time", "horizon_hours"])
    rebuilt = []
    for row in raw3.itertuples(index=False, name=None):
        date, issue, *values = row
        date = pd.Timestamp(date).strftime("%Y-%m-%d")
        for horizon, value in enumerate(values, 1):
            rebuilt.append((date, issue, horizon, float(value)))
    rebuilt.sort(key=lambda x: (x[0], x[1], x[2]))
    assert [(d, t, h) for d, t, h, _ in rebuilt] == list(zip(normalized.issue_date, normalized.issue_time, normalized.horizon_hours))
    close([v for *_, v in rebuilt], normalized.pv_forecast_kw)

    for stem in ("plan", "adjusted", "charge", "discharge", "emergency"):
        assert frame[stem + "_kw"].min() >= -TOL
        close(frame[stem + "_kwh"], frame[stem + "_kw"] * DT)
    assert frame[["charge_kw", "discharge_kw"]].max().max() <= 5000 + TOL
    assert not ((frame.charge_kw > TOL) & (frame.discharge_kw > TOL)).any()
    assert frame[["soc_start_kwh", "soc_end_kwh"]].min().min() >= 1200 - TOL
    assert frame[["soc_start_kwh", "soc_end_kwh"]].max().max() <= 10800 + TOL
    close(frame.soc_start_kwh.iloc[0], 6000)
    close(frame.soc_start_kwh.to_numpy()[1:], frame.soc_end_kwh.to_numpy()[:-1])
    close(frame.soc_end_kwh - frame.soc_start_kwh, 0.9 * frame.charge_kw * DT - frame.discharge_kw * DT / 0.9)
    close(
        frame.adjusted_kw + frame.emergency_kw + frame.pv_actual_kw + frame.discharge_kw,
        frame.load_actual_kw + frame.charge_kw + frame.surplus_kw,
    )
    # No forecast issued after a slot starts may affect that slot.
    assert (frame.forecast_issue_hour <= np.floor(frame.slot / 6)).all()
    for hour, lower, upper in ((0, 0, 36), (6, 36, 72), (12, 72, 108), (18, 108, 144)):
        block = delivered[(delivered.slot >= lower) & (delivered.slot < upper)]
        assert (block.forecast_issue_hour == hour).all()
    diff = frame.adjusted_kw - frame.plan_kw
    adjustment = np.where(diff >= 0, 1.5 * frame.price_yuan_per_kwh * diff * DT, 0.5 * frame.price_yuan_per_kwh * diff * DT)
    close(frame.plan_cost_yuan, frame.price_yuan_per_kwh * frame.plan_kwh)
    close(frame.adjustment_cost_yuan, adjustment)
    close(frame.emergency_cost_yuan, 5 * frame.price_yuan_per_kwh * frame.emergency_kwh)
    close(frame.total_cost_yuan, frame.plan_cost_yuan + frame.adjustment_cost_yuan + frame.emergency_cost_yuan)
    for column in (
        "plan_kwh", "adjusted_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh",
        "plan_cost_yuan", "adjustment_cost_yuan", "ordinary_settlement_cost_yuan",
        "emergency_cost_yuan", "total_cost_yuan",
    ):
        close(delivered[column].sum(), summary[column])
        close(delivered.groupby("date")[column].sum().to_numpy(), daily[column])

    expected_events = []
    for date, group in delivered.groupby("date", sort=True):
        positive = group.emergency_kwh.to_numpy() > 1e-8
        edges = np.diff(np.r_[False, positive, False].astype(int))
        for start, stop in zip(np.where(edges == 1)[0], np.where(edges == -1)[0]):
            expected_events.append((date, f"{time_label(start)}-{time_label(stop)}", group.emergency_kwh.iloc[start:stop].sum()))
    assert list(zip(events.date, events.interval)) == [(d, t) for d, t, _ in expected_events]
    close(events.emergency_kwh, [v for _, _, v in expected_events])

    book = load_workbook(OUT / "result3.xlsx", data_only=True)
    template = load_workbook(ROOT / "题目/附件/附件5/result3.xlsx", data_only=True)
    assert book.sheetnames == template.sheetnames
    plan_sheet, adjusted_sheet = book["计划购电量"], book["调整购电量"]
    assert plan_sheet.max_row == adjusted_sheet.max_row == 335
    for sheet in (plan_sheet, adjusted_sheet):
        assert [sheet.cell(1, c).value for c in range(1, 148)] == [template[sheet.title].cell(1, c).value for c in range(1, 148)]
    storage = book["充放电量"]
    assert storage.max_row == 2005
    for i, (date, group) in enumerate(delivered.groupby("date", sort=True)):
        assert pd.Timestamp(plan_sheet.cell(i + 2, 1).value).strftime("%Y-%m-%d") == date
        close([plan_sheet.cell(i + 2, c).value for c in range(2, 146)], group.plan_kwh)
        close([plan_sheet.cell(i + 2, 146).value, plan_sheet.cell(i + 2, 147).value], [group.plan_kwh.sum(), group.plan_cost_yuan.sum()])
        close([adjusted_sheet.cell(i + 2, c).value for c in range(2, 146)], group.adjusted_kwh)
        close([adjusted_sheet.cell(i + 2, 146).value, adjusted_sheet.cell(i + 2, 147).value], [group.adjusted_kwh.sum(), group.ordinary_settlement_cost_yuan.sum()])
        row = 2 + 6 * i
        for col, field in ((3, "charge_kwh"), (4, "discharge_kwh")):
            close([storage.cell(row + block, col).value for block in range(6)], group[field].to_numpy().reshape(6, 24).sum(axis=1))
        close([storage.cell(row, 6).value, storage.cell(row + 1, 6).value], [group.soc_start_kwh.iloc[0], group.soc_end_kwh.iloc[-1]])
    displayed, workbook_events, current_date = set(), [], None
    for date, label, energy in book["紧急购电量"].iter_rows(min_row=2, max_col=3, values_only=True):
        if date is not None:
            current_date = pd.Timestamp(date).strftime("%Y-%m-%d")
            displayed.add(current_date)
        if label == "无":
            close(energy, 0)
        elif label is not None:
            workbook_events.append((current_date, label, energy))
    assert displayed == set(delivered.date)
    assert [(d, t) for d, t, _ in workbook_events] == [(d, t) for d, t, _ in expected_events]
    close([v for _, _, v in workbook_events], [v for _, _, v in expected_events])
    assert not any(cell.data_type == "e" for sheet in book for row in sheet for cell in row)
    print(json.dumps({
        "verified": True, "delivered_days": 334, "delivered_slots": 48096,
        "emergency_events": len(expected_events), "total_cost_yuan": summary["total_cost_yuan"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
